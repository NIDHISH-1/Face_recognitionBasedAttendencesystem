import cv2
import os
from deepface import DeepFace
from datetime import datetime
from openpyxl import Workbook, load_workbook

# -------------------------------
# PATH SETUP
# -------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
db_path = os.path.join(BASE_DIR, "dataset")
excel_file = os.path.join(BASE_DIR, "attendance.xlsx")

# -------------------------------
# CREATE EXCEL FILE IF NOT EXISTS
# -------------------------------
if not os.path.exists(excel_file):
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws.append(["Name", "Date", "Time"])
    wb.save(excel_file)

# Load workbook
wb = load_workbook(excel_file)
ws = wb.active

# -------------------------------
# FACE DETECTOR (OpenCV)
# -------------------------------
face_cascade = cv2.CascadeClassifier(
    cv2.data.haarcascades + "haarcascade_frontalface_default.xml"
)

# -------------------------------
# TRACK ATTENDANCE (NO DUPLICATE)
# -------------------------------
marked = set()

# -------------------------------
# START CAMERA
# -------------------------------
cap = cv2.VideoCapture(0)

print("Starting Attendance System... Press 'q' to exit")

while True:
    ret, frame = cap.read()
    if not ret:
        break

    # Resize for better performance
    frame = cv2.resize(frame, (640, 480))

    gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
    faces = face_cascade.detectMultiScale(gray, 1.3, 5)

    for (x, y, w, h) in faces:
        face_img = frame[y:y+h, x:x+w]

        try:
            result = DeepFace.find(
                img_path=face_img,
                db_path=db_path,
                enforce_detection=False
            )

            if len(result) > 0 and len(result[0]) > 0:
                identity = result[0].iloc[0]['identity']
                name = os.path.basename(os.path.dirname(identity))

                # MARK ATTENDANCE
                if name not in marked:
                    now = datetime.now()
                    date = now.strftime("%Y-%m-%d")
                    time = now.strftime("%H:%M:%S")

                    ws.append([name, date, time])
                    wb.save(excel_file)

                    marked.add(name)
                    print(f"Attendance marked for {name}")

                # DRAW GREEN BOX
                cv2.rectangle(frame, (x, y), (x+w, y+h), (0, 255, 0), 2)
                cv2.putText(frame, name, (x, y-10),
                            cv2.FONT_HERSHEY_SIMPLEX, 0.9, (0, 255, 0), 2)

            else:
                # UNKNOWN FACE
                cv2.rectangle(frame, (x, y), (x+w, y+h), (0, 0, 255), 2)
                cv2.putText(frame, "Unknown", (x, y-10),
                            cv2.FONT_HERSHEY_SIMPLEX, 0.9, (0, 0, 255), 2)

        except Exception as e:
            print("Error:", e)

    cv2.imshow("Face Recognition Attendance", frame)

    if cv2.waitKey(1) & 0xFF == ord('q'):
        break

cap.release()
cv2.destroyAllWindows()

# -------------------------------
# CREATE SUMMARY SHEET
# -------------------------------
if "Summary" not in wb.sheetnames:
    summary = wb.create_sheet("Summary")
    summary.append(["Name", "Total उपस्थित"])

    names = set(row[0].value for row in ws.iter_rows(min_row=2))

    for n in names:
        count = sum(1 for row in ws.iter_rows(min_row=2) if row[0].value == n)
        summary.append([n, count])

    wb.save(excel_file)